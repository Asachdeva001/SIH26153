from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR

OUT = "SIH26_Predictive_Cyber_Defense.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

W, H = prs.slide_width, prs.slide_height
NAVY = RGBColor(7, 40, 84)
BLUE = RGBColor(24, 112, 188)
GREEN = RGBColor(0, 122, 55)
ORANGE = RGBColor(255, 153, 51)
INK = RGBColor(20, 20, 20)
PALE_BLUE = RGBColor(232, 242, 250)
PALE_GREEN = RGBColor(239, 249, 241)
PALE_ORANGE = RGBColor(255, 247, 237)
GREY = RGBColor(95, 105, 112)
WHITE = RGBColor(255, 255, 255)
RED = RGBColor(191, 34, 34)


def set_fill(shape, color):
    shape.fill.solid(); shape.fill.fore_color.rgb = color


def set_line(shape, color, width=1):
    shape.line.color.rgb = color; shape.line.width = Pt(width)


def textbox(slide, x, y, w, h, text, size=14, color=INK, bold=False,
            font="Arial", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
            margin=0.04, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(margin); tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin); tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = font; r.font.size = Pt(size); r.font.bold = bold
    r.font.italic = italic; r.font.color.rgb = color
    return box


def richbox(slide, x, y, w, h, runs, size=12, color=INK, font="Arial", margin=0.08):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(margin); tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin); tf.margin_bottom = Inches(margin)
    p = tf.paragraphs[0]
    for text, col, bold in runs:
        r = p.add_run(); r.text = text; r.font.name = font; r.font.size = Pt(size)
        r.font.color.rgb = col; r.font.bold = bold
    return box


def rect(slide, x, y, w, h, fill=WHITE, line=BLUE, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(shape, fill); set_line(shape, line, 1)
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def title(slide, text, number):
    textbox(slide, 0.32, 0.05, 2.0, 0.42, "SIH26", 20, NAVY, True, "Arial Narrow")
    textbox(slide, 0.32, 0.38, 1.95, 0.22, "Predictive SOC", 9, BLUE, True, "Arial Narrow")
    textbox(slide, 10.9, 0.10, 2.1, 0.26, "SMART INDIA", 10, NAVY, True, "Arial Narrow", PP_ALIGN.RIGHT)
    textbox(slide, 10.7, 0.35, 2.3, 0.26, "HACKATHON 2025", 10, GREY, True, "Arial Narrow", PP_ALIGN.RIGHT)
    textbox(slide, 2.0, 0.12, 8.7, 0.55, text, 26, BLUE, True, "Arial Narrow", PP_ALIGN.CENTER)
    rect(slide, 0, 7.18, 13.333, 0.32, BLUE, BLUE)
    textbox(slide, 0.35, 7.20, 8.5, 0.22, "@ SIH26 - PREDICTIVE CYBER DEFENSE", 10, WHITE, True, "Arial Narrow")
    textbox(slide, 12.65, 7.20, 0.35, 0.22, str(number), 11, WHITE, True, "Arial Narrow", PP_ALIGN.RIGHT)


def bullet_list(slide, x, y, w, h, items, size=11, color=INK, gap=0.03):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(0.06); tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02); tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "• " + item; p.font.name = "Arial"; p.font.size = Pt(size)
        p.font.color.rgb = color; p.space_after = Pt(gap * 72)
    return box


def panel(slide, x, y, w, h, heading, fill=PALE_BLUE, accent=BLUE):
    rect(slide, x, y, w, h, fill, accent, True)
    textbox(slide, x+0.10, y+0.06, w-0.20, 0.28, heading.upper(), 13, accent, True, "Arial Narrow")


def arrow(slide, x1, y1, x2, y2, color=BLUE, width=2):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color; line.line.width = Pt(width); line.line.end_arrowhead = True
    return line

# Slide 1
s = prs.slides.add_slide(prs.slide_layouts[6]); title(s, "PREDICTIVE CYBER DEFENSE", 1)
textbox(s, 0.55, 1.05, 6.0, 0.42, "Problem Statement: Predict the attack before impact", 19, INK, True)
items = [
    "Static IDS tools report what happened; SOC teams still need to predict what happens next.",
    "Multi-stage campaigns evolve across time windows, assets and ATT&CK stages.",
    "Analysts need one explainable view that connects risk, attack stage, asset criticality and action."
]
bullet_list(s, 0.65, 1.65, 5.75, 1.55, items, 14)
panel(s, 0.55, 3.48, 5.75, 2.55, "Our Solution", PALE_GREEN, GREEN)
bullet_list(s, 0.72, 3.95, 5.35, 1.75, [
    "A temporal World Model that learns network-state transitions.",
    "K-step forecasting of future risk and predicted traffic state.",
    "MITRE ATT&CK mapping, GradientShap explanation and SOC playbook prioritization.",
    "Local, containerized inference for sensitive CII and air-gapped environments."
], 13)
rect(s, 7.0, 1.20, 5.55, 4.80, PALE_BLUE, BLUE, True)
textbox(s, 7.30, 1.50, 4.95, 0.34, "FROM TELEMETRY TO DECISION", 17, NAVY, True, "Arial Narrow", PP_ALIGN.CENTER)
for i, (label, col) in enumerate([("NETWORK\nTELEMETRY", BLUE), ("WORLD MODEL\nFORECAST", ORANGE), ("SOC ACTION\nPRIORITY", GREEN)]):
    x = 7.35 + i*1.75
    rect(s, x, 2.45, 1.40, 1.15, col, col, True)
    textbox(s, x+0.08, 2.72, 1.24, 0.55, label, 13, WHITE, True, "Arial Narrow", PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    if i < 2: arrow(s, x+1.42, 3.02, x+1.70, 3.02, NAVY, 2)
textbox(s, 7.35, 4.15, 4.90, 1.15, "A working FastAPI + Next.js prototype that turns network behavior into an auditable, explainable response workflow.", 15, NAVY, True, "Arial", PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
textbox(s, 0.60, 6.45, 12.0, 0.40, "Theme: Software   |   Focus: National cyber resilience, predictive SOC operations, critical infrastructure", 12, GREEN, True, "Arial", PP_ALIGN.CENTER)

# Slide 2
s = prs.slides.add_slide(prs.slide_layouts[6]); title(s, "SOLUTION & ARCHITECTURE", 2)
panel(s, 0.35, 0.95, 5.4, 1.48, "Problem Existing", PALE_ORANGE, ORANGE)
bullet_list(s, 0.52, 1.38, 5.05, 0.88, ["Alerts are point-in-time and noisy.", "Risk is disconnected from asset importance.", "Analyst decisions are difficult to explain and audit."], 11)
panel(s, 0.35, 2.58, 5.4, 1.78, "Proposed Solution", PALE_GREEN, GREEN)
bullet_list(s, 0.52, 3.00, 5.05, 1.12, ["CSV/PCAP telemetry becomes 10-second state windows.", "LSTM World Model rolls the state forward K steps.", "ATT&CK + XAI + SOC priority convert prediction into action."], 11)
panel(s, 0.35, 4.52, 5.4, 1.52, "Unique Value Proposition", PALE_BLUE, BLUE)
bullet_list(s, 0.52, 4.93, 5.05, 0.90, ["Forecast lead time, not just detection.", "One analyst workflow from signal to response.", "Local inference with browser audit history."], 11)
rect(s, 6.05, 0.95, 6.90, 5.10, WHITE, NAVY, True)
textbox(s, 6.25, 1.08, 6.45, 0.28, "ARCHITECTURE", 14, NAVY, True, "Arial Narrow", PP_ALIGN.CENTER)
flow = [("CSV / PCAP\nUPLOAD", 6.35, 2.00, BLUE), ("PARSER\n10s WINDOWS", 8.00, 2.00, ORANGE), ("PYTORCH\nWORLD MODEL", 9.65, 2.00, NAVY), ("K-STEP\nFORECAST", 11.30, 2.00, GREEN)]
for label, x, y, col in flow:
    rect(s, x, y, 1.40, 0.82, col, col, True)
    textbox(s, x+0.04, y+0.16, 1.32, 0.45, label, 10, WHITE, True, "Arial Narrow", PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
for x in [7.78, 9.43, 11.08]: arrow(s, x, 2.41, x+0.20, 2.41, NAVY, 2)
for label, x, y, col in [("MITRE\nSTAGE", 7.10, 3.65, BLUE), ("XAI\nDRIVERS", 8.90, 3.65, ORANGE), ("SOC\nPRIORITY", 10.70, 3.65, GREEN)]:
    rect(s, x, y, 1.45, 0.85, PALE_BLUE if col == BLUE else PALE_GREEN if col == GREEN else PALE_ORANGE, col, True)
    textbox(s, x+0.05, y+0.18, 1.35, 0.43, label, 10, col, True, "Arial Narrow", PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
arrow(s, 10.30, 2.84, 7.85, 3.63, BLUE, 1.5); arrow(s, 10.30, 2.84, 9.62, 3.63, ORANGE, 1.5); arrow(s, 11.95, 2.84, 11.40, 3.63, GREEN, 1.5)
rect(s, 6.50, 5.10, 5.90, 0.60, PALE_BLUE, BLUE, True)
textbox(s, 6.65, 5.25, 5.60, 0.28, "Next.js SOC dashboard  ↔  FastAPI REST API  ↔  PyTorch inference", 12, NAVY, True, "Arial", PP_ALIGN.CENTER)

# Slide 3
s = prs.slides.add_slide(prs.slide_layouts[6]); title(s, "TECHNICAL APPROACH", 3)
textbox(s, 0.38, 0.88, 12.5, 0.32, "The system learns how a network state changes over time, then simulates the next K states.", 14, NAVY, True, "Arial", PP_ALIGN.CENTER)
# left pipeline
rect(s, 0.38, 1.36, 7.20, 4.92, PALE_BLUE, BLUE, True)
textbox(s, 0.62, 1.52, 6.70, 0.28, "MODEL PIPELINE", 14, BLUE, True, "Arial Narrow", PP_ALIGN.CENTER)
steps = [("1", "Feature engineering", "25 telemetry features per 10-second window"), ("2", "Temporal memory", "2-layer LSTM, hidden dimension 64, history length 8"), ("3", "Dual prediction heads", "Next-state decoder + attack risk head"), ("4", "Autoregressive rollout", "Predicted state feeds the next future step"), ("5", "Decision layer", "ATT&CK stage, XAI drivers, asset-aware priority")]
for i, (n, h, d) in enumerate(steps):
    y = 1.98 + i*0.78
    rect(s, 0.72, y, 0.43, 0.43, ORANGE if i < 4 else GREEN, ORANGE if i < 4 else GREEN, True)
    textbox(s, 0.72, y+0.07, 0.43, 0.22, n, 12, WHITE, True, "Arial", PP_ALIGN.CENTER)
    textbox(s, 1.35, y-0.02, 2.10, 0.28, h, 11, NAVY, True)
    textbox(s, 3.45, y-0.02, 3.75, 0.36, d, 10, INK)
    if i < 4: arrow(s, 0.93, y+0.44, 0.93, y+0.72, BLUE, 1.5)
# right boxes
panel(s, 7.88, 1.36, 5.05, 1.42, "Input Representation", PALE_ORANGE, ORANGE)
bullet_list(s, 8.08, 1.82, 4.65, 0.70, ["Flow count, packets, bytes, TCP flags, IAT, TTL, IP/port diversity, scan and high-port signals."], 10)
panel(s, 7.88, 2.98, 5.05, 1.42, "Explainability", PALE_GREEN, GREEN)
bullet_list(s, 8.08, 3.44, 4.65, 0.70, ["Captum GradientShap returns top feature attributions and a human-readable SOC narrative."], 10)
panel(s, 7.88, 4.60, 5.05, 1.68, "Training & Evaluation", PALE_BLUE, BLUE)
bullet_list(s, 8.08, 5.03, 4.65, 1.02, ["Standardized states; MSE state loss + weighted BCE risk loss.", "Chronological campaign grouping and Logistic Regression baseline.", "K-horizon evaluation includes F1, precision, recall, FPR and lead time."], 10)

# Slide 4
s = prs.slides.add_slide(prs.slide_layouts[6]); title(s, "FEASIBILITY & VIABILITY", 4)
panel(s, 0.42, 0.92, 6.05, 3.12, "Feasibility", PALE_BLUE, BLUE)
textbox(s, 0.68, 1.38, 5.55, 0.34, "Technical", 12, NAVY, True)
bullet_list(s, 0.78, 1.75, 5.35, 0.72, ["FastAPI + PyTorch + Next.js stack is implemented and containerized.", "Runs on CPU or CUDA; model loads at service startup."], 10)
textbox(s, 0.68, 2.58, 5.55, 0.34, "Operational", 12, NAVY, True)
bullet_list(s, 0.78, 2.95, 5.35, 0.72, ["Docker Compose exposes the dashboard on 3000 and API on 8000.", "Local inference avoids request-time cloud dependency for sensitive networks."], 10)
panel(s, 6.82, 0.92, 6.08, 3.12, "Viability", PALE_GREEN, GREEN)
textbox(s, 7.08, 1.38, 5.55, 0.34, "Who benefits", 12, GREEN, True)
bullet_list(s, 7.18, 1.75, 5.40, 0.72, ["SOC analysts receive lead time and a recommended playbook.", "CII operators get asset-criticality-aware prioritization.", "Security leadership gets an auditable decision trail."], 10)
textbox(s, 7.08, 2.58, 5.55, 0.34, "Deployment model", 12, GREEN, True)
bullet_list(s, 7.18, 2.95, 5.40, 0.72, ["Cloud can train; protected on-prem environments can run inference.", "Architecture leaves room for SOAR, EDR and firewall adapters."], 10)
textbox(s, 0.42, 4.42, 12.3, 0.30, "CURRENT CHALLENGES & MITIGATION", 15, NAVY, True, "Arial Narrow", PP_ALIGN.CENTER)
panel(s, 0.42, 4.86, 3.88, 1.55, "Data realism", PALE_ORANGE, ORANGE)
bullet_list(s, 0.62, 5.28, 3.48, 0.82, ["Validate artifact provenance.", "Disclose packet-level defaults.", "Retrain on verified real traffic."], 10)
panel(s, 4.72, 4.86, 3.88, 1.55, "Production hardening", PALE_ORANGE, ORANGE)
bullet_list(s, 4.92, 5.28, 3.48, 0.82, ["Version model manifests.", "Add async upload jobs.", "Add rate limiting and auth."], 10)
panel(s, 9.02, 4.86, 3.88, 1.55, "Action integration", PALE_ORANGE, ORANGE)
bullet_list(s, 9.22, 5.28, 3.48, 0.82, ["Connect real EDR/SOAR.", "Replace simulator actions.", "Keep human approval in loop."], 10)

# Slide 5
s = prs.slides.add_slide(prs.slide_layouts[6]); title(s, "IMPACT & BENEFITS", 5)
textbox(s, 0.45, 0.88, 12.4, 0.30, "From a noisy alert queue to a forecast-led, explainable SOC decision.", 15, NAVY, True, "Arial", PP_ALIGN.CENTER)
# benefits left
panel(s, 0.42, 1.35, 5.15, 4.80, "Operational impact", PALE_BLUE, BLUE)
bullet_list(s, 0.70, 1.84, 4.58, 3.45, [
    "Lead time: forecast future risk across configurable K steps.",
    "Clarity: translate model output into MITRE ATT&CK stages.",
    "Prioritization: combine forecast risk, stage severity and asset criticality.",
    "Explainability: show which network features drive the prediction.",
    "Governance: preserve alert acknowledgement and isolation actions in browser audit history.",
    "Resilience: keep inference local for sensitive government and CII environments."
], 12)
# metrics right
panel(s, 5.90, 1.35, 6.98, 2.35, "Evidence snapshot", PALE_GREEN, GREEN)
textbox(s, 6.20, 1.83, 2.00, 0.72, "26", 30, GREEN, True, "Arial Narrow", PP_ALIGN.CENTER)
textbox(s, 8.22, 1.93, 1.80, 0.42, "forecast lead\nwindows", 12, NAVY, True, "Arial Narrow", PP_ALIGN.CENTER)
textbox(s, 10.15, 1.83, 1.90, 0.72, "0.6313", 25, BLUE, True, "Arial Narrow", PP_ALIGN.CENTER)
textbox(s, 11.98, 1.93, 0.68, 0.42, "World\nModel F1", 9, NAVY, True, "Arial Narrow", PP_ALIGN.CENTER)
textbox(s, 6.20, 2.82, 6.25, 0.48, "Checked-in benchmark artifact: World Model recall 0.6212 vs baseline 0.6085; results require provenance verification before production claims.", 10, RED, True, "Arial", PP_ALIGN.CENTER)
# workflow modes
panel(s, 5.90, 3.98, 6.98, 2.17, "How the SOC workflow works", PALE_ORANGE, ORANGE)
workflow = [("OBSERVE", "Telemetry window\n+ current risk", BLUE), ("FORECAST", "K-step risk\ntrajectory", ORANGE), ("ACT", "Stage + asset\n+ playbook", GREEN)]
for i, (h, d, col) in enumerate(workflow):
    x = 6.20 + i*2.18
    rect(s, x, 4.62, 1.73, 0.86, col, col, True)
    textbox(s, x+0.04, 4.75, 1.65, 0.22, h, 12, WHITE, True, "Arial Narrow", PP_ALIGN.CENTER)
    textbox(s, x+0.06, 5.00, 1.61, 0.32, d, 9, WHITE, False, "Arial", PP_ALIGN.CENTER)
    if i < 2: arrow(s, x+1.76, 5.05, x+2.08, 5.05, NAVY, 2)

# Slide 6
s = prs.slides.add_slide(prs.slide_layouts[6]); title(s, "RESEARCH, EVIDENCE & ROADMAP", 6)
panel(s, 0.42, 0.90, 5.95, 2.30, "Implemented evidence", PALE_BLUE, BLUE)
bullet_list(s, 0.68, 1.38, 5.40, 1.45, [
    "FastAPI routes: scenario, upload, forecast, MITRE, XAI, SOC, benchmark and report.",
    "Eight dashboard views: Forecast, ATT&CK, XAI, K-step, SOC, Benchmark, Topology and Audit.",
    "Automated tests cover forecasting, parsing, routes, evaluation, XAI and prioritization.",
    "Docker Compose provides reproducible local deployment."
], 10)
panel(s, 6.68, 0.90, 6.20, 2.30, "Research basis", PALE_GREEN, GREEN)
bullet_list(s, 6.95, 1.38, 5.65, 1.45, [
    "World-model framing: learn transition dynamics P(S(t+1) | S(t)).",
    "Temporal sequence learning with a two-layer LSTM.",
    "Captum GradientShap for local feature attribution.",
    "MITRE ATT&CK vocabulary for analyst-facing interpretation."
], 10)
textbox(s, 0.42, 3.52, 12.4, 0.30, "WINNING PATH: PROVE, HARDEN, INTEGRATE", 15, NAVY, True, "Arial Narrow", PP_ALIGN.CENTER)
road = [("01", "Prove", "Re-run on verified CIC-IDS-2018 data; publish a versioned benchmark and model manifest.", BLUE), ("02", "Harden", "Add authentication, rate limits, asynchronous upload jobs and production observability.", ORANGE), ("03", "Integrate", "Connect approved EDR, SOAR, SIEM and firewall adapters with human-in-the-loop controls.", GREEN)]
for i, (n, h, d, col) in enumerate(road):
    x = 0.55 + i*4.25
    rect(s, x, 4.05, 3.85, 1.58, WHITE, col, True)
    rect(s, x+0.15, 4.22, 0.60, 0.60, col, col, True)
    textbox(s, x+0.15, 4.39, 0.60, 0.20, n, 11, WHITE, True, "Arial Narrow", PP_ALIGN.CENTER)
    textbox(s, x+0.90, 4.18, 2.65, 0.27, h.upper(), 13, col, True, "Arial Narrow")
    textbox(s, x+0.90, 4.55, 2.68, 0.75, d, 10, INK)
textbox(s, 0.55, 6.05, 12.0, 0.43, "References: repository README and audit reports | backend tests | benchmark_report.json | MITRE ATT&CK | Captum / PyTorch documentation", 10, GREY, False, "Arial", PP_ALIGN.CENTER)
textbox(s, 0.55, 6.52, 12.0, 0.28, "The differentiator: forecast the next state, explain why, prioritize what matters, and preserve the decision.", 13, GREEN, True, "Arial", PP_ALIGN.CENTER)

prs.save(OUT)
print(OUT)
